"""Offline contract, native OOXML, adversarial package, and CLI acceptance tests.

No renderer/validator mocks: mutations start with a real rendered package and
change one independently observable property. A rendering failure deliberately
blocks those checks rather than substituting a more permissive fixture.
"""

import json
import os
import socket
import subprocess
import sys
import unicodedata
import warnings
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pydantic import ValidationError

from mini_deerflow.demo.phase2_pptx_offline import demo_spec
from mini_deerflow.pptx import DeckSpec, render_deck, validate_pptx
from mini_deerflow.pptx.design import THEMES
from mini_deerflow.pptx.renderer import LayoutOverflowError
from mini_deerflow.pptx.schemas import SlideSpec
from mini_deerflow.pptx.validator import (
    MAX_FILE_BYTES,
    MAX_UNPACKED_BYTES,
    PptxValidationError,
    ValidationReport,
)

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/package/2006/relationships",
}
KINDS = {
    "title",
    "bullets",
    "comparison",
    "metric",
    "steps",
    "chart",
    "table",
    "conclusion",
}
ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("PPTX tests must never access the network")

    monkeypatch.setattr(socket.socket, "connect", forbidden)
    monkeypatch.setattr(socket, "create_connection", forbidden)


@pytest.fixture(scope="module")
def spec():
    return demo_spec()


@pytest.fixture(scope="module")
def rendered(spec):
    return render_deck(spec)


def payload():
    return demo_spec().model_dump(mode="json")


def put(data, path, value):
    for key in path[:-1]:
        data = data[key]
    data[path[-1]] = value


def xml_mutation(blob, member, change):
    with ZipFile(BytesIO(blob)) as archive:
        root = ET.fromstring(archive.read(member))
    change(root)
    return repack(
        blob, {member: ET.tostring(root, encoding="utf-8", xml_declaration=True)}
    )


def repack(blob, replacements=None, *, remove=(), extras=None):
    replacements = replacements or {}
    output = BytesIO()
    with ZipFile(BytesIO(blob)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            if entry.filename not in remove:
                target.writestr(
                    entry.filename, replacements.get(entry.filename, source.read(entry))
                )
        for name, data in (extras or {}).items():
            target.writestr(name, data)
    return output.getvalue()


def edit_presentation(blob, change):
    prs = Presentation(BytesIO(blob))
    change(prs)
    output = BytesIO()
    prs.save(output)
    return output.getvalue()


def named(slide, name):
    return next(shape for shape in slide.shapes if shape.name == name)


def reject(blob, spec, message=None, **kwargs):
    with pytest.raises(PptxValidationError, match=message):
        validate_pptx(blob, spec, **kwargs)


TEXT_FIELDS = [
    (("title",), 80),
    (("slides", 0, "title"), 80),
    (("slides", 0, "content", "subtitle"), 180),
    (("slides", 1, "content", "items", 0), 180),
    (("slides", 2, "content", "left", "heading"), 36),
    (("slides", 2, "content", "right", "items", 0), 180),
    (("slides", 3, "content", "value"), 36),
    (("slides", 3, "content", "label"), 80),
    (("slides", 3, "content", "insight"), 180),
    (("slides", 4, "content", "items", 0), 80),
    (("slides", 5, "content", "categories", 0), 36),
    (("slides", 5, "content", "series", 0, "name"), 36),
    (("slides", 5, "content", "caption"), 180),
    (("slides", 6, "content", "headers", 0), 48),
    (("slides", 6, "content", "rows", 0, 0), 48),
    (("slides", 6, "content", "caption"), 180),
    (("slides", 7, "content", "items", 0), 180),
    (("slides", 7, "content", "sources", 0), 180),
]


@pytest.mark.parametrize("path,limit", TEXT_FIELDS)
def test_text_accepts_exact_budget_and_rejects_next_character(path, limit):
    data = payload()
    put(data, path, "ế" * limit)
    assert DeckSpec.model_validate(data)
    put(data, path, "ế" * (limit + 1))
    with pytest.raises(ValidationError):
        DeckSpec.model_validate(data)


@pytest.mark.parametrize(
    "text",
    [
        "",
        " ",
        "\u2003",
        "\t",
        "\n",
        "x\x00y",
        "x\x1fy",
        "x\x7fy",
        "x\x85y",
        chr(0xD800),
        chr(0xDFFF),
        "\ufffe",
        "\uffff",
    ],
)
@pytest.mark.parametrize(
    "path",
    [
        ("title",),
        ("slides", 0, "content", "subtitle"),
        ("slides", 6, "content", "rows", 0, 0),
    ],
)
def test_unsafe_or_blank_text_rejected_at_nested_boundaries(text, path):
    data = payload()
    put(data, path, text)
    with pytest.raises(ValidationError):
        DeckSpec.model_validate(data)


@pytest.mark.parametrize(
    "text",
    [
        "Nghiên cứu tiếng Việt — Đổi mới",
        "<tag> & \"literal\" 'text'",
        "漢字 🙂",
        unicodedata.normalize("NFD", "Tiếng Việt"),
    ],
)
def test_safe_unicode_and_literal_markup_are_preserved(text):
    data = payload()
    data["slides"][0]["content"]["subtitle"] = text
    actual = DeckSpec.model_validate(data)
    assert actual.slides[0].content.subtitle == text
    assert DeckSpec.model_validate_json(actual.model_dump_json()) == actual


@pytest.mark.parametrize(
    "key,value",
    [
        ("aspect_ratio", "4:3"),
        ("aspect_ratio", "16:10"),
        ("theme", "unknown"),
        ("theme", "Dark"),
    ],
)
def test_closed_aspect_ratio_and_theme(key, value):
    data = payload()
    data[key] = value
    with pytest.raises(ValidationError):
        DeckSpec.model_validate(data)


@pytest.mark.parametrize(
    "id_", ["", "1slide", "Upper", "has space", "has_under", "é", "a" * 49, "a\n"]
)
def test_slide_id_rejects_invalid_identifiers(id_):
    data = payload()
    data["slides"][0]["id"] = id_
    with pytest.raises(ValidationError):
        DeckSpec.model_validate(data)


def test_slide_id_and_deck_count_boundaries_and_uniqueness():
    slide = payload()["slides"][0]
    assert SlideSpec.model_validate({**slide, "id": "a" + "0-" * 23 + "z"})
    for count in (1, 40):
        data = {
            "title": "Deck",
            "slides": [{**slide, "id": f"s-{i}"} for i in range(count)],
        }
        assert len(DeckSpec.model_validate(data).slides) == count
    for slides in ([], [{**slide, "id": f"s-{i}"} for i in range(41)], [slide, slide]):
        with pytest.raises(ValidationError):
            DeckSpec.model_validate({"title": "Deck", "slides": slides})


@pytest.mark.parametrize(
    "path",
    [
        (),
        ("slides", 0),
        ("slides", 0, "content"),
        ("slides", 2, "content", "left"),
        ("slides", 5, "content", "series", 0),
    ],
)
def test_extra_fields_forbidden_recursively(path):
    data = payload()
    target = data
    for key in path:
        target = target[key]
    target["unexpected"] = "must not be silently discarded"
    with pytest.raises(ValidationError, match="extra_forbidden"):
        DeckSpec.model_validate(data)


@pytest.mark.parametrize(
    "index,field,low,high",
    [(1, "items", 1, 6), (4, "items", 2, 5), (7, "items", 1, 3), (7, "sources", 1, 3)],
)
def test_item_count_boundaries(index, field, low, high):
    data = payload()
    for count in (low, high):
        data["slides"][index]["content"][field] = ["Nội dung"] * count
        assert DeckSpec.model_validate(data)
    for count in (low - 1, high + 1):
        data["slides"][index]["content"][field] = ["Nội dung"] * count
        with pytest.raises(ValidationError):
            DeckSpec.model_validate(data)


@pytest.mark.parametrize("count,valid", [(0, False), (1, True), (4, True), (5, False)])
def test_comparison_column_counts(count, valid):
    data = payload()
    data["slides"][2]["content"]["left"]["items"] = ["Ý chính"] * count
    if valid:
        assert DeckSpec.model_validate(data)
    else:
        with pytest.raises(ValidationError):
            DeckSpec.model_validate(data)


@pytest.mark.parametrize(
    "value",
    [
        float("nan"),
        float("inf"),
        -float("inf"),
        1_000_000_001,
        -1_000_000_001,
        True,
        None,
        "12",
    ],
)
@pytest.mark.parametrize("location", ["chart", "table"])
def test_numeric_values_must_be_finite_bounded_and_not_coerced(value, location):
    data = payload()
    if location == "chart":
        data["slides"][5]["content"]["series"][0]["values"][0] = value
    else:
        data["slides"][6]["content"]["rows"][0][1] = value
        # Text cells are legal; strings are not a numeric field in a table.
        if isinstance(value, str):
            assert DeckSpec.model_validate(data).slides[6].content.rows[0][1] == value
            return
    with pytest.raises(ValidationError):
        DeckSpec.model_validate(data)


@pytest.mark.parametrize("value", [-1_000_000_000, 0, 1_000_000_000, -0.125])
def test_numeric_endpoints_are_supported(value):
    data = payload()
    data["slides"][5]["content"]["series"][0]["values"][0] = value
    data["slides"][6]["content"]["rows"][0][1] = value
    actual = DeckSpec.model_validate(data)
    assert actual.slides[5].content.series[0].values[0] == value
    assert actual.slides[6].content.rows[0][1] == value


@pytest.mark.parametrize("categories,series", [(1, 1), (8, 3)])
def test_chart_bounds_accept_endpoints(categories, series):
    data = payload()
    chart = data["slides"][5]["content"]
    chart["categories"] = [f"Mục {i}" for i in range(categories)]
    chart["series"] = [
        {"name": f"Chuỗi {i}", "values": list(range(categories))} for i in range(series)
    ]
    assert DeckSpec.model_validate(data)


@pytest.mark.parametrize(
    "mutation",
    [
        "empty-categories",
        "too-many-categories",
        "empty-series",
        "too-many-series",
        "empty-values",
        "too-many-values",
        "ragged",
        "duplicate-category",
        "duplicate-series",
        "unsupported-type",
        "unknown-kind",
    ],
)
def test_chart_contract_rejects_invalid_structure(mutation):
    data = payload()
    chart = data["slides"][5]["content"]
    if mutation == "empty-categories":
        chart["categories"] = []
    elif mutation == "too-many-categories":
        chart["categories"] = [str(i) for i in range(9)]
        for series in chart["series"]:
            series["values"] = list(range(9))
    elif mutation == "empty-series":
        chart["series"] = []
    elif mutation == "too-many-series":
        chart["series"] = [{"name": str(i), "values": [1, 2, 3]} for i in range(4)]
    elif mutation in {"empty-values", "too-many-values", "ragged"}:
        chart["series"][0]["values"] = {
            "empty-values": [],
            "too-many-values": list(range(9)),
            "ragged": [1, 2],
        }[mutation]
    elif mutation == "duplicate-category":
        chart["categories"][1] = chart["categories"][0]
    elif mutation == "duplicate-series":
        chart["series"][1]["name"] = chart["series"][0]["name"]
    elif mutation == "unsupported-type":
        chart["chart_type"] = "pie"
    else:
        chart["kind"] = "html"
    with pytest.raises(ValidationError):
        DeckSpec.model_validate(data)


@pytest.mark.parametrize(
    "columns,rows,valid",
    [
        (1, 1, True),
        (5, 8, True),
        (0, 1, False),
        (6, 1, False),
        (1, 0, False),
        (1, 9, False),
    ],
)
def test_table_dimension_bounds(columns, rows, valid):
    data = payload()
    table = data["slides"][6]["content"]
    table["headers"] = [f"C{i}" for i in range(columns)]
    table["rows"] = [["ô"] * columns for _ in range(rows)]
    if valid:
        assert DeckSpec.model_validate(data)
    else:
        with pytest.raises(ValidationError):
            DeckSpec.model_validate(data)


@pytest.mark.parametrize("row", [[], ["short"], ["too", "many", "cells", "here"]])
def test_table_requires_rectangular_rows(row):
    data = payload()
    data["slides"][6]["content"]["rows"][0] = row
    with pytest.raises(ValidationError, match="table rows"):
        DeckSpec.model_validate(data)


def test_contract_is_frozen_and_renderer_revalidates_model_copy(spec):
    with pytest.raises(ValidationError, match="frozen"):
        spec.theme = "dark"
    unsafe = spec.model_copy(update={"theme": "not-a-theme"})
    with pytest.raises(ValidationError):
        render_deck(unsafe)
    forged = DeckSpec.model_construct(
        title="unsafe", slides=(), theme="light", aspect_ratio="16:9"
    )
    with pytest.raises(ValidationError):
        render_deck(forged)


def test_density_overflow_is_reported_for_schema_valid_content():
    dense = DeckSpec.model_validate(
        {
            "title": "Dense",
            "slides": [
                {
                    "id": "dense",
                    "title": "Dense",
                    "content": {"kind": "bullets", "items": ["漢" * 180] * 6},
                }
            ],
        }
    )
    sparse_data = dense.model_dump(mode="json")
    sparse_data["slides"][0]["content"]["items"] = ["Short readable item"]
    # Establish that the same layout can render before attributing a rejection
    # to dense content; an unconditional renderer failure is not a passing test.
    assert render_deck(DeckSpec.model_validate(sparse_data)).startswith(b"PK")
    with pytest.raises(LayoutOverflowError, match="density"):
        render_deck(dense)


@pytest.mark.parametrize("theme_name", ["light", "dark", "academic"])
def test_eight_slides_native_objects_unicode_and_theme_consistency(
    theme_name, tmp_path
):
    expected = demo_spec(theme_name)
    blob = render_deck(expected)
    assert isinstance(blob, bytes) and blob.startswith(b"PK")
    report = validate_pptx(blob, expected, min_layouts=8)
    assert isinstance(report, ValidationReport) and report.valid
    assert report.slide_count == 8 and set(report.layouts) == KINDS
    assert report.native_objects == {
        "text_boxes": 33,
        "shapes": 14,
        "tables": 1,
        "charts": 1,
    }
    assert report.geometry_checked == 49
    assert report.file_bytes == len(blob)
    assert report.visual_rendering_verified is False
    path = tmp_path / "tiếng-Việt.pptx"
    path.write_bytes(blob)
    assert validate_pptx(path, expected, min_layouts=8) == report
    assert report.to_dict()["slide_count"] == 8
    prs = Presentation(BytesIO(blob))
    assert (prs.slide_width, prs.slide_height) == (11704320, 6583680)
    assert prs.core_properties.title == expected.title
    theme = THEMES[theme_name]
    counts_by_layout = {
        "title": (3, 1, 0, 0),
        "bullets": (3, 1, 0, 0),
        "comparison": (6, 3, 0, 0),
        "metric": (5, 1, 0, 0),
        "steps": (6, 5, 0, 0),
        "chart": (3, 1, 0, 1),
        "table": (3, 1, 1, 0),
        "conclusion": (4, 1, 0, 0),
    }
    for slide, wanted, inspection in zip(
        prs.slides, expected.slides, report.slides, strict=True
    ):
        assert slide.name == f"{wanted.id}:{wanted.content.kind}"
        assert named(slide, "title").text == wanted.title
        assert (inspection.id, inspection.layout, inspection.title) == (
            wanted.id,
            wanted.content.kind,
            wanted.title,
        )
        counts = counts_by_layout[wanted.content.kind]
        assert (
            inspection.text_boxes,
            inspection.shapes,
            inspection.tables,
            inspection.charts,
        ) == counts
        assert inspection.geometry_checked == len(slide.shapes) == sum(counts)
        assert str(slide.background.fill.fore_color.rgb) == theme.background
        assert str(named(slide, "accent").fill.fore_color.rgb) == theme.accent
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width
            assert shape.top + shape.height <= prs.slide_height
            if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
                color = (
                    theme.accent
                    if shape.name == "value"
                    else theme.muted
                    if shape.name == "sources"
                    else theme.foreground
                )
                for paragraph in shape.text_frame.paragraphs:
                    assert paragraph.font.name == theme.font
                    assert str(paragraph.font.color.rgb) == color
            elif shape.has_table:
                for ri, row in enumerate(shape.table.rows):
                    for cell in row.cells:
                        assert cell.text_frame.paragraphs[0].font.name == theme.font
                        assert (
                            str(cell.text_frame.paragraphs[0].font.color.rgb)
                            == theme.foreground
                        )
                        assert str(cell.fill.fore_color.rgb) == (
                            theme.surface if ri == 0 else theme.background
                        )
            elif shape.has_chart:
                assert shape.chart.font.name == theme.font
                assert str(shape.chart.font.color.rgb) == theme.foreground
                assert [
                    str(s.format.fill.fore_color.rgb) for s in shape.chart.series
                ] == [theme.accent, theme.muted]
            else:
                assert str(shape.fill.fore_color.rgb) == (
                    theme.accent if shape.name == "accent" else theme.surface
                )
    assert named(prs.slides[0], "subtitle").text == expected.slides[0].content.subtitle
    assert named(prs.slides[2], "left-heading").text == "Ảnh tĩnh"
    chart = named(prs.slides[5], "chart").chart
    assert (
        tuple(c.label for c in chart.plots[0].categories)
        == expected.slides[5].content.categories
    )
    assert tuple((s.name, tuple(s.values)) for s in chart.series) == tuple(
        (s.name, s.values) for s in expected.slides[5].content.series
    )
    table = named(prs.slides[6], "table").table
    assert tuple(tuple(c.text for c in row.cells) for row in table.rows) == tuple(
        tuple(str(v) for v in row)
        for row in (
            expected.slides[6].content.headers,
            *expected.slides[6].content.rows,
        )
    )


@pytest.mark.parametrize(
    "chart_type,enum",
    [
        ("column", XL_CHART_TYPE.COLUMN_CLUSTERED),
        ("bar", XL_CHART_TYPE.BAR_CLUSTERED),
        ("line", XL_CHART_TYPE.LINE),
    ],
)
def test_chart_types_render_as_editable_native_charts(chart_type, enum):
    data = payload()
    data["slides"] = [data["slides"][5]]
    data["slides"][0]["content"]["chart_type"] = chart_type
    expected = DeckSpec.model_validate(data)
    blob = render_deck(expected)
    assert validate_pptx(blob, expected).native_objects["charts"] == 1
    assert (
        named(Presentation(BytesIO(blob)).slides[0], "chart").chart.chart_type == enum
    )


def semantic_parts(blob):
    """Ignore ZIP clocks/order; include nested workbook XML and binary members."""
    with ZipFile(BytesIO(blob)) as archive:
        result = {}
        for name in sorted(archive.namelist()):
            data = archive.read(name)
            if name.endswith((".xml", ".rels")):
                result[name] = ET.canonicalize(data.decode("utf-8"))
            elif name.endswith(".xlsx"):
                result[name] = semantic_parts(data)
            else:
                result[name] = data
        return result


@pytest.mark.parametrize("theme", ["light", "dark", "academic"])
def test_repeat_render_has_identical_semantic_xml(theme):
    expected = demo_spec(theme)
    first, second = render_deck(expected), render_deck(expected)
    assert semantic_parts(first) == semantic_parts(second)
    assert validate_pptx(first, expected).valid
    assert validate_pptx(second, expected).valid


@pytest.mark.parametrize("blob", [b"", b"not a zip", b"PK\x03\x04truncated"])
def test_corrupt_or_empty_input_raises_public_validation_error(blob, spec):
    reject(blob, spec)


def test_crc_corruption_is_rejected_before_xml_parsing(spec):
    output = BytesIO()
    # Stored compression makes the exact payload location unambiguous. Keep the
    # central directory CRC unchanged while flipping one member byte.
    marker = b"<presentation>unique-crc-marker</presentation>"
    with ZipFile(output, "w") as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr("ppt/presentation.xml", marker)
    damaged = bytearray(output.getvalue())
    location = damaged.index(marker)
    damaged[location + 1] ^= 1
    reject(bytes(damaged), spec, "CRC")


@pytest.mark.parametrize(
    "member", ["ppt/presentation.xml", "ppt/slides/slide1.xml", "ppt/charts/chart1.xml"]
)
def test_corrupted_xml_is_rejected(rendered, spec, member):
    reject(repack(rendered, {member: b"<unclosed>"}), spec, "invalid PPTX")


def test_dtd_entity_expansion_is_rejected_as_public_error(spec):
    xml = b'<!DOCTYPE root [<!ENTITY x "unsafe">]><root>&x;</root>'
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr("ppt/presentation.xml", xml)
    reject(output.getvalue(), spec)


@pytest.mark.parametrize(
    "target",
    ["deck-title", "slide-title", "content", "identity", "page", "canvas", "count"],
)
def test_tampered_identity_text_and_canvas_are_rejected(rendered, spec, target):
    def change(prs):
        if target == "deck-title":
            prs.core_properties.title = "Altered deck"
        elif target == "canvas":
            prs.slide_width += 100
        elif target == "count":
            prs.slides.add_slide(prs.slide_layouts[6])
        elif target == "identity":
            prs.slides[0].name = "changed:title"
        else:
            name = {"slide-title": "title", "content": "subtitle", "page": "page"}[
                target
            ]
            named(prs.slides[0], name).text = "Nội dung bị thay đổi"

    reject(edit_presentation(rendered, change), spec)


@pytest.mark.parametrize(
    "axis,value",
    [
        ("left", -1),
        ("top", -1),
        ("left", 11704320),
        ("top", 6583680),
        ("width", 0),
        ("height", 0),
        ("rotation", 5),
    ],
)
def test_off_canvas_zero_size_and_rotation_rejected(rendered, spec, axis, value):
    bad = edit_presentation(
        rendered, lambda prs: setattr(named(prs.slides[0], "subtitle"), axis, value)
    )
    reject(bad, spec, "geometry")


@pytest.mark.parametrize(
    "slide_index,name",
    [
        (0, "subtitle"),
        (0, "accent"),
        (2, "left-panel"),
        (4, "step-panel-0"),
        (5, "chart"),
        (6, "table"),
    ],
)
def test_missing_native_objects_are_rejected(rendered, spec, slide_index, name):
    def change(prs):
        element = named(prs.slides[slide_index], name)._element
        element.getparent().remove(element)

    reject(edit_presentation(rendered, change), spec)


def test_duplicate_native_object_names_are_rejected(rendered, spec):
    def change(prs):
        slide = prs.slides[0]
        slide.shapes._spTree.insert_element_before(
            deepcopy(named(slide, "subtitle")._element), "p:extLst"
        )

    reject(edit_presentation(rendered, change), spec, "duplicate object names")


def test_changed_table_data_is_rejected(rendered, spec):
    def change(prs):
        named(prs.slides[6], "table").table.cell(1, 1).text = "999"

    reject(edit_presentation(rendered, change), spec, "table cell")


@pytest.mark.parametrize("xpath", [".//c:numCache/c:pt/c:v", ".//c:strCache/c:pt/c:v"])
def test_changed_chart_cache_is_rejected(rendered, spec, xpath):
    def change(root):
        node = root.find(xpath, NS)
        assert node is not None, "mutation must reach a real cache cell"
        node.text = "999" if "numCache" in xpath else "Tampered"

    reject(xml_mutation(rendered, "ppt/charts/chart1.xml", change), spec)


def test_changed_embedded_workbook_is_rejected_independently_of_cache(rendered, spec):
    with ZipFile(BytesIO(rendered)) as archive:
        (member,) = [
            n
            for n in archive.namelist()
            if n.startswith("ppt/embeddings/") and n.endswith(".xlsx")
        ]
        workbook = archive.read(member)
        chart_before = archive.read("ppt/charts/chart1.xml")

    def change(root):
        value = root.find(".//s:c[@r='B2']/s:v", NS)
        assert value is not None
        value.text = "999"

    workbook = xml_mutation(workbook, "xl/worksheets/sheet1.xml", change)
    bad = repack(rendered, {member: workbook})
    with ZipFile(BytesIO(bad)) as archive:
        assert archive.read("ppt/charts/chart1.xml") == chart_before
    reject(bad, spec, "embedded chart workbook data mismatch")


def test_layout_coverage_threshold_is_enforced(rendered, spec):
    assert validate_pptx(rendered, spec, min_layouts=8).valid
    reject(rendered, spec, "layout coverage", min_layouts=9)


def test_one_slide_requires_only_one_layout_by_default():
    data = payload()
    data["slides"] = data["slides"][:1]
    expected = DeckSpec.model_validate(data)
    blob = render_deck(expected)
    assert validate_pptx(blob, expected).layouts == ("title",)
    reject(blob, expected, "layout coverage", min_layouts=2)


@pytest.mark.parametrize("member", ["../escape.xml", "/absolute.xml", "bad\\path.xml"])
def test_unsafe_zip_member_paths_are_rejected(rendered, spec, member):
    entry = ZipInfo(member)
    # ZipInfo construction normalizes platform separators on Windows. Restore
    # the literal unsafe name so writestr cannot silently sanitize this fixture.
    entry.filename = member
    bad = repack(rendered, extras={entry: b"<root/>"})
    assert member.encode("utf-8") in bad
    with ZipFile(BytesIO(bad)) as archive:
        assert member in [info.orig_filename for info in archive.infolist()]
    reject(bad, spec, "unsafe ZIP member path")


def test_duplicate_zip_members_rejected(rendered, spec):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        bad = repack(rendered, extras={"ppt/presentation.xml": b"<root/>"})
    reject(bad, spec, "duplicate ZIP members")


@pytest.mark.parametrize(
    "member", ["[Content_Types].xml", "ppt/presentation.xml", "ppt/charts/chart1.xml"]
)
def test_missing_parts_or_relationship_targets_rejected(rendered, spec, member):
    reject(repack(rendered, remove={member}), spec)


@pytest.mark.parametrize("mutation", ["external", "duplicate-id", "missing-target"])
def test_invalid_relationships_rejected(rendered, spec, mutation):
    def change(root):
        if mutation == "external":
            root[0].set("TargetMode", "External")
            root[0].set("Target", "https://example.invalid/resource")
        elif mutation == "duplicate-id":
            root[1].set("Id", root[0].get("Id"))
        else:
            root[0].set("Target", "nonexistent.xml")

    reject(xml_mutation(rendered, "ppt/_rels/presentation.xml.rels", change), spec)


def test_compressed_and_path_file_limits_without_rendering(spec, tmp_path):
    oversized = b"x" * (MAX_FILE_BYTES + 1)
    reject(oversized, spec, "file size limit")
    path = tmp_path / "oversized.pptx"
    path.write_bytes(oversized)
    reject(path, spec, "file size limit")
    reject(tmp_path / "missing.pptx", spec, "invalid PPTX")


def test_unpacked_zip_limit_without_rendering(spec):
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("padding.bin", b"0" * (MAX_UNPACKED_BYTES + 1))
    assert len(output.getvalue()) < MAX_FILE_BYTES
    reject(output.getvalue(), spec, "ZIP expansion limit")


def test_zip_member_count_limit_without_rendering(spec):
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for index in range(2001):
            archive.writestr(f"part-{index}.bin", b"")
    reject(output.getvalue(), spec, "too many ZIP members")


@pytest.mark.parametrize("theme", ["light", "dark", "academic"])
def test_offline_demo_subprocess_real_output_and_refuses_overwrite(tmp_path, theme):
    # The guard lives only in this temporary subprocess, and permits no socket
    # connection even if an accidental future dependency tries to make one.
    bootstrap = (
        "import runpy,socket; "
        "exec('def blocked(*a, **k):\\n raise AssertionError(\"network forbidden\")'); "
        "socket.socket.connect=blocked; socket.create_connection=blocked; "
        "runpy.run_module('mini_deerflow.demo.phase2_pptx_offline',run_name='__main__')"
    )
    output_dir = tmp_path / "kết quả"
    command = [
        sys.executable,
        "-c",
        bootstrap,
        "--output-dir",
        str(output_dir),
        "--theme",
        theme,
    ]
    env = {**os.environ, "UV_OFFLINE": "1", "PYTHONUTF8": "1"}
    first = subprocess.run(
        command,
        cwd=ROOT,
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=60,
        check=False,
    )
    assert first.returncode == 0, first.stderr
    report = json.loads(first.stdout)
    path = output_dir / "phase2-offline.pptx"
    assert Path(report["file"]) == path.resolve()
    assert report["valid"] and report["slide_count"] == 8
    assert set(report["layouts"]) == KINDS
    original = path.read_bytes()
    assert original.startswith(b"PK")
    assert validate_pptx(path, demo_spec(theme), min_layouts=8).valid
    second = subprocess.run(
        command,
        cwd=ROOT,
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=60,
        check=False,
    )
    assert second.returncode != 0
    assert "FileExistsError" in second.stderr
    assert path.read_bytes() == original
    assert list(output_dir.iterdir()) == [path]
