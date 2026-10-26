"""Credential-free visual fixture and evidence-contract regressions.

Only the dedicated Linux viewer job asserts that LibreOffice actually rendered
the decks; these unit tests must never claim Office visual approval.
"""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import pytest
from pptx import Presentation

from mini_deerflow.pptx import render_deck, validate_pptx
from scripts import visual_acceptance as visual


@pytest.mark.parametrize("lang", visual.LANGUAGES)
@pytest.mark.parametrize("theme", visual.THEMES)
def test_matrix_deck_has_all_native_layouts_and_editable_chart_table(lang, theme):
    spec = visual.make_spec(lang, theme)
    assert [slide.content.kind for slide in spec.slides] == list(visual.LAYOUTS)
    assert spec.theme == theme
    assert len(spec.slides) == 8
    source = spec.slides[7].content.sources[0]
    assert 100 <= len(source) <= 180
    pptx = render_deck(spec)
    report = validate_pptx(pptx, spec, min_layouts=8)
    assert report.valid and report.slide_count == 8
    prs = Presentation(BytesIO(pptx))
    assert len(prs.slides) == 8
    assert any(shape.has_chart for shape in prs.slides[5].shapes)
    assert any(shape.has_table for shape in prs.slides[6].shapes)


def test_vietnamese_diacritics_and_english_sample_have_independent_copy():
    vi = visual.make_spec("vi", "academic")
    en = visual.make_spec("en", "academic")
    assert "ă â ê ô ơ ư đ" in vi.slides[1].content.items[0]
    assert "Ă Â Ê Ô Ơ Ư Đ" in vi.slides[1].content.items[0]
    assert "Editable" in en.slides[3].content.label
    assert vi.slides[1].content.items != en.slides[1].content.items
    assert "Tháng một" in vi.slides[5].content.categories
    assert "January" in en.slides[5].content.categories


def test_generation_only_is_explicitly_blocked_not_accepted(tmp_path: Path):
    report = visual.run(tmp_path)
    manifest = json.loads((tmp_path / "manifest.json").read_text(encoding="utf-8"))
    assert report["overall_status"] == "BLOCKED_VIEWER_NOT_RUN"
    assert manifest["visual_reviewer_status"] == "UNVERIFIED"
    assert manifest["powerpoint_visual_status"] == "UNVERIFIED"
    assert manifest["manual_edit_resave_status"] == "UNVERIFIED"
    assert len(manifest["matrix"]) == 6
    assert all(x["viewer_render"] == "NOT_RUN" for x in manifest["matrix"])
    assert all(x["rendered_pages"] == 0 for x in manifest["matrix"])
    assert all(x == "UNVERIFIED" for x in manifest["manual_checklist"].values())
    assert len(list((tmp_path / "pptx").glob("*.pptx"))) == 6
    assert not (tmp_path / "screenshots").exists()


def test_viewer_unavailable_is_blocked_and_never_faked(tmp_path, monkeypatch):
    monkeypatch.setattr(visual.shutil, "which", lambda command: None)
    monkeypatch.setattr(
        visual.subprocess, "run",
        lambda *args, **kwargs: pytest.fail("no viewer executable may be invoked"),
    )
    report = visual.run(tmp_path, render=True)
    assert report["overall_status"] == "BLOCKED_VIEWER_UNAVAILABLE"
    assert all(x["viewer_render"] == "NOT_RUN" for x in report["matrix"])
    assert not (tmp_path / "screenshots").exists()
    manifest = json.loads((tmp_path / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["overall_status"] == "BLOCKED_VIEWER_UNAVAILABLE"


def test_png_evidence_rejects_bad_or_tiny_images(tmp_path: Path):
    path = tmp_path / "slide.png"
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\0" * 8
                     + (1400).to_bytes(4, "big") + (788).to_bytes(4, "big"))
    assert visual.png_dimensions(path) == (1400, 788)
    path.write_bytes(b"not an actual png")
    with pytest.raises(ValueError, match="not a PNG"):
        visual.png_dimensions(path)
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\0" * 8
                     + (100).to_bytes(4, "big") + (80).to_bytes(4, "big"))
    with pytest.raises(ValueError, match="too small"):
        visual.png_dimensions(path)


def test_page_count_fail_closed_and_pdf_text_is_only_a_hint():
    assert visual._pages("Title: Example\nPages:      8\nAuthor: fixture") == 8
    with pytest.raises(ValueError, match="page count unknown"):
        visual._pages("empty")
    assert visual._pdf_text_coverage("Mini DeerFlow; editable; Source", "en")
    assert not visual._pdf_text_coverage("Mini DeerFlow only", "en")
    assert not visual._pdf_text_coverage("", "vi")


def test_native_edit_roundtrip_preserves_text_and_chart_object():
    """Programmatic OOXML edit, explicitly NOT a real Office manual-edit test."""
    spec = visual.make_spec("en", "light")
    source = render_deck(spec)
    presentation = Presentation(BytesIO(source))
    table = next(shape.table for shape in presentation.slides[6].shapes if shape.has_table)
    table.cell(1, 0).text = "Edited and saved"
    output = BytesIO()
    presentation.save(output)
    reopened = Presentation(BytesIO(output.getvalue()))
    updated_table = next(
        shape.table for shape in reopened.slides[6].shapes if shape.has_table
    )
    assert updated_table.cell(1, 0).text == "Edited and saved"
    assert any(shape.has_chart for shape in reopened.slides[5].shapes)


def test_invalid_fixture_parameter_fails_without_provider():
    with pytest.raises(ValueError, match="unsupported"):
        visual.make_spec("xx", "light")
    with pytest.raises(ValueError, match="unsupported"):
        visual.make_spec("vi", "night")
