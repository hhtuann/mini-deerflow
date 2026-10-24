"""Issue10 Task2.2: real PPTX publication, hostile storage, and HTTP contracts.

Only local filesystem/SQLite failures are injected. Rendering, validation,
conversation persistence, the offline runtime, and FastAPI remain real.
"""

import asyncio
import hashlib
import json
import os
import sqlite3
import subprocess
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import aiosqlite
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from mini_deerflow.api import create_app
from mini_deerflow.conversation import SQLiteConversationRepository
from mini_deerflow.demo.offline_scenario import OfflineDemoBackend
from mini_deerflow.demo.phase2_pptx_offline import demo_spec
from mini_deerflow.demo.service import DemoRuntimeService, RunDemoCommand
from mini_deerflow.pptx import DeckSpec, render_deck, validate_pptx
from mini_deerflow.pptx.artifacts import (
    MAX_SPEC_BYTES,
    PPTX_MIME,
    ArtifactConflictError,
    ArtifactError,
    ArtifactIntegrityError,
    ArtifactNotFoundError,
    PptxArtifactStore,
)
from mini_deerflow.sandbox import SessionSandboxResolver
from mini_deerflow.workspace import (
    Workspace,
    WorkspaceLimitError,
    WorkspacePathError,
    _filesystem_path,
)

THREAD, TURN = "artifact-thread", "artifact-turn"


def digest(data):
    return hashlib.sha256(data).hexdigest()


def canonical(spec):
    return json.dumps(
        spec.model_dump(mode="json"),
        ensure_ascii=True,
        sort_keys=True,
        separators=(",", ":"),
    )


def repack(data, *, year=2024, remove=()):
    """Change ZIP timestamps deterministically without changing OOXML content."""
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w") as target:
        for name in source.namelist():
            if name not in remove:
                info = ZipInfo(name, date_time=(year, 1, 1, 0, 0, 0))
                info.compress_type = ZIP_DEFLATED
                target.writestr(info, source.read(name))
    return output.getvalue()


@pytest.fixture(scope="module")
def deck():
    spec = demo_spec()
    assert len(spec.slides) == 8
    data = render_deck(spec)
    validate_pptx(data, spec)
    return spec, data


@pytest.fixture(scope="module")
def other_deck(deck):
    payload = deck[0].model_dump(mode="json")
    payload["slides"][0]["title"] = "A separately owned presentation"
    spec = DeckSpec.model_validate(payload)
    return spec, render_deck(spec)


async def seed(root, thread=THREAD, turn=TURN, *, completed=True):
    """Create ownership through the public real conversation repository."""
    repo = SQLiteConversationRepository(root / "checkpoints.sqlite")
    await repo.create_conversation_with_turn(
        thread_id=thread,
        turn_id=turn,
        user_message="Publish a bounded local presentation.",
        session_tool_call_limit=20,
        reserved_tool_calls=5,
    )
    if completed:
        await complete(repo, thread, turn)
    return repo


async def complete(repo, thread=THREAD, turn=TURN):
    await repo.complete_turn(
        thread_id=thread,
        turn_id=turn,
        assistant_response="Presentation evidence is complete.",
        total_tool_calls=0,
    )


def sql(root, statement, parameters=()):
    with sqlite3.connect(root / "checkpoints.sqlite") as connection:
        connection.execute("PRAGMA foreign_keys=ON")
        return connection.execute(statement, parameters).fetchall()


def artifact_path(root, record):
    # Independently verify the documented sandbox and ownership-addressed layout.
    identity = json.dumps(
        (record.mode, record.thread_id, record.turn_id), separators=(",", ":")
    )
    sandbox = SessionSandboxResolver(root / "workspaces").resolve(record.thread_id)
    return _filesystem_path(
        sandbox / "presentations" / digest(identity.encode()) / f"{record.sha256}.pptx"
    )


def filesystem_matches(root, pattern):
    """Walk from an extended root, keeping comparisons in the same namespace.

    Prefixing only discovered files is too late: ordinary Windows rglob can
    silently miss descendants beyond MAX_PATH. Return I/O-ready paths so crash
    recovery assertions can also read the discovered files directly.
    """
    return sorted(
        _filesystem_path(path) for path in _filesystem_path(root).rglob(pattern)
    )


def safe_error(error, root):
    text = str(error).replace("\\\\", "\\")
    assert str(root) not in text
    assert "Traceback" not in text


def test_publication_snapshot_schema_foreign_keys_and_fresh_store(tmp_path, deck):
    async def scenario():
        spec, data = deck
        await seed(tmp_path)
        store = PptxArtifactStore(tmp_path, mode="offline")
        record = await store.publish(THREAD, TURN, data, spec)
        assert record.mode == "offline"
        assert (record.thread_id, record.turn_id, record.artifact_type) == (
            THREAD,
            TURN,
            "pptx",
        )
        assert record.byte_size == len(data)
        assert record.sha256 == digest(data)
        assert record.spec_json == canonical(spec)
        assert record.spec_sha256 == digest(record.spec_json.encode())
        assert DeckSpec.model_validate_json(record.spec_json) == spec
        assert artifact_path(tmp_path, record).read_bytes() == data
        fresh = PptxArtifactStore(tmp_path, mode="offline")
        assert await fresh.retrieve(THREAD, TURN) == (record, data)
        assert len(Presentation(BytesIO(data)).slides) == 8
        async with fresh._connection() as connection:
            assert (await (await connection.execute("PRAGMA foreign_keys")).fetchone())[
                0
            ] == 1
        foreign_keys = sql(tmp_path, "PRAGMA foreign_key_list(pptx_artifacts)")
        assert {(row[2], row[3], row[4], row[6]) for row in foreign_keys} == {
            ("turns", "thread_id", "thread_id", "CASCADE"),
            ("turns", "turn_id", "turn_id", "CASCADE"),
        }
        columns = sql(tmp_path, "PRAGMA table_info(pptx_artifacts)")
        assert {row[1] for row in columns if row[5]} == {"mode", "thread_id", "turn_id"}
        with pytest.raises(sqlite3.IntegrityError):
            sql(tmp_path, "UPDATE pptx_artifacts SET turn_id='nonexistent'")
        assert sql(tmp_path, "PRAGMA foreign_key_check") == []
        assert sql(tmp_path, "SELECT count(*) FROM turns") == [(1,)]

    asyncio.run(scenario())


def test_canonical_retry_concurrent_publish_and_conflict(tmp_path, deck, other_deck):
    async def scenario():
        spec, data = deck
        await seed(tmp_path)
        stores = [PptxArtifactStore(tmp_path, mode="offline") for _ in range(4)]
        timestamped = repack(data, year=2025)
        assert timestamped != data
        validate_pptx(timestamped, spec)
        first = await stores[0].publish(THREAD, TURN, data, spec)
        # Reordered input JSON produces the same canonical specification snapshot.
        equivalent = DeckSpec.model_validate(
            dict(reversed(list(spec.model_dump().items())))
        )
        results = await asyncio.gather(
            *(store.publish(THREAD, TURN, timestamped, equivalent) for store in stores)
        )
        assert results == [first] * 4
        assert await stores[-1].retrieve(THREAD, TURN) == (first, data)
        with pytest.raises(ArtifactConflictError):
            await stores[0].publish(THREAD, TURN, other_deck[1], other_deck[0])
        assert sql(tmp_path, "SELECT count(*) FROM pptx_artifacts") == [(1,)]
        assert filesystem_matches(tmp_path, "*.pptx") == [
            artifact_path(tmp_path, first)
        ]

    asyncio.run(scenario())


@pytest.mark.parametrize("different_spec", [False, True])
def test_first_publication_race_is_immutable(
    tmp_path, deck, other_deck, different_spec
):
    async def scenario():
        await seed(tmp_path)
        contenders = (
            deck,
            other_deck if different_spec else (deck[0], repack(deck[1])),
        )
        results = await asyncio.gather(
            *(
                PptxArtifactStore(tmp_path, mode="offline").publish(
                    THREAD, TURN, data, spec
                )
                for spec, data in contenders
            ),
            return_exceptions=True,
        )
        winners = [item for item in results if not isinstance(item, Exception)]
        assert len(winners) == (1 if different_spec else 2), results
        if different_spec:
            assert sum(isinstance(item, ArtifactConflictError) for item in results) == 1
        else:
            assert winners[0] == winners[1]
        record, data = await PptxArtifactStore(tmp_path, mode="offline").retrieve(
            THREAD, TURN
        )
        assert record in winners
        assert digest(data) == record.sha256
        assert len(filesystem_matches(tmp_path, "*.pptx")) == 1

    asyncio.run(scenario())


def test_turn_thread_and_mode_isolation_with_same_identifiers(
    tmp_path, deck, other_deck
):
    async def scenario():
        for mode, root, presentation in (
            ("offline", tmp_path, deck),
            ("live", tmp_path / "live", other_deck),
        ):
            repo = await seed(root)
            store = PptxArtifactStore(root, mode=mode)
            first = await store.publish(THREAD, TURN, presentation[1], presentation[0])
            await repo.reserve_follow_up(
                thread_id=THREAD,
                turn_id="next-turn",
                user_message="A distinct deck.",
                reserved_tool_calls=5,
            )
            await complete(repo, THREAD, "next-turn")
            second = await store.publish(THREAD, "next-turn", deck[1], deck[0])
            await seed(root, "other-thread", TURN)
            third = await store.publish(
                "other-thread", TURN, other_deck[1], other_deck[0]
            )
            assert (
                len({artifact_path(root, item) for item in (first, second, third)}) == 3
            )
            assert (await store.retrieve(THREAD, TURN))[1] == presentation[1]
            assert (await store.retrieve(THREAD, "next-turn"))[1] == deck[1]
            assert (await store.retrieve("other-thread", TURN))[1] == other_deck[1]
        with pytest.raises(ArtifactIntegrityError):
            await PptxArtifactStore(tmp_path, mode="live").retrieve(THREAD, TURN)

    asyncio.run(scenario())


@pytest.mark.parametrize("status", ["running", "interrupted", "failed", "completed"])
def test_turn_status_controls_publish_and_download(tmp_path, deck, status):
    async def scenario():
        repo = await seed(tmp_path, completed=False)
        store = PptxArtifactStore(tmp_path, mode="offline")
        record = await store.publish(THREAD, TURN, deck[1], deck[0])
        if status == "completed":
            await complete(repo)
            assert await store.retrieve(THREAD, TURN) == (record, deck[1])
        else:
            if status == "interrupted":
                await repo.mark_turn_interrupted(thread_id=THREAD, turn_id=TURN)
            elif status == "failed":
                await repo.mark_turn_failed(
                    thread_id=THREAD, turn_id=TURN, total_tool_calls=0
                )
            with pytest.raises(ArtifactNotFoundError):
                await store.retrieve(THREAD, TURN)
            if status != "running":
                with pytest.raises(ArtifactNotFoundError):
                    await store.publish(THREAD, TURN, deck[1], deck[0])

    asyncio.run(scenario())


def test_unknown_and_wrong_pairs_never_read_files(tmp_path, deck, monkeypatch):
    async def scenario():
        await seed(tmp_path)
        await seed(tmp_path, "other-thread", "other-turn")
        store = PptxArtifactStore(tmp_path, mode="offline")
        await store.publish(THREAD, TURN, deck[1], deck[0])

        def forbidden(*args, **kwargs):
            pytest.fail("Unknown ownership must be rejected before workspace access")

        monkeypatch.setattr(Workspace, "read_bytes", forbidden)
        monkeypatch.setattr(Workspace, "write_bytes_exclusive", forbidden)
        for thread, turn in (
            ("missing", TURN),
            (THREAD, "missing"),
            ("other-thread", TURN),
            (THREAD, "other-turn"),
            ("other-thread", "other-turn"),
        ):
            with pytest.raises(ArtifactNotFoundError):
                await store.retrieve(thread, turn)
        for thread, turn in (("missing", TURN), ("other-thread", TURN)):
            with pytest.raises(ArtifactNotFoundError):
                await store.publish(thread, turn, deck[1], deck[0])

    asyncio.run(scenario())


@pytest.mark.parametrize(
    "field,value",
    [
        ("sha256", "0" * 64),
        ("sha256", "z" * 64),
        ("byte_size", 1),
        ("spec_sha256", "f" * 64),
        ("spec_json", "{}"),
        ("mode", "live"),
    ],
)
def test_metadata_tampering_fails_closed(tmp_path, deck, field, value):
    async def scenario():
        await seed(tmp_path)
        store = PptxArtifactStore(tmp_path, mode="offline")
        record = await store.publish(THREAD, TURN, deck[1], deck[0])
        sql(tmp_path, f"UPDATE pptx_artifacts SET {field}=?", (value,))
        expected = ArtifactNotFoundError if field == "mode" else ArtifactIntegrityError
        with pytest.raises(expected) as error:
            await store.retrieve(THREAD, TURN)
        safe_error(error.value, tmp_path)
        assert artifact_path(tmp_path, record).read_bytes() == deck[1]

    asyncio.run(scenario())


@pytest.mark.parametrize(
    "damage", ["missing", "truncated", "same-size", "invalid-package", "invalid-spec"]
)
def test_read_revalidates_bytes_package_and_spec_even_with_updated_hashes(
    tmp_path, deck, damage
):
    async def scenario():
        await seed(tmp_path)
        store = PptxArtifactStore(tmp_path, mode="offline")
        record = await store.publish(THREAD, TURN, deck[1], deck[0])
        path = artifact_path(tmp_path, record)
        if damage == "missing":
            path.unlink()
        elif damage == "invalid-spec":
            snapshot = '{"title":"invalid contract"}'
            sql(
                tmp_path,
                "UPDATE pptx_artifacts SET spec_json=?, spec_sha256=?",
                (snapshot, digest(snapshot.encode())),
            )
        elif damage == "invalid-package":
            data = repack(deck[1], remove=("ppt/slides/slide8.xml",))
            # An attacker updates both the hash-addressed filename and all byte metadata.
            path.unlink()
            path.with_name(f"{digest(data)}.pptx").write_bytes(data)
            sql(
                tmp_path,
                "UPDATE pptx_artifacts SET sha256=?, byte_size=?",
                (digest(data), len(data)),
            )
        elif damage == "truncated":
            path.write_bytes(deck[1][:100])
        else:
            path.write_bytes(bytes([deck[1][0] ^ 1]) + deck[1][1:])
        with pytest.raises(ArtifactIntegrityError) as error:
            await store.retrieve(THREAD, TURN)
        safe_error(error.value, tmp_path)

    asyncio.run(scenario())


@pytest.mark.parametrize("failure", ["insert", "commit"])
def test_transaction_failure_leaves_recoverable_file_without_metadata(
    tmp_path, deck, monkeypatch, failure
):
    async def scenario():
        await seed(tmp_path)
        store = PptxArtifactStore(tmp_path, mode="offline")
        original_execute = aiosqlite.Connection.execute
        original_commit = aiosqlite.Connection.commit
        inserted = False

        async def execute(connection, statement, parameters=None):
            nonlocal inserted
            if statement.startswith("INSERT INTO pptx_artifacts VALUES"):
                inserted = True
                if failure == "insert":
                    raise sqlite3.OperationalError(f"injected failure {tmp_path}")
            return await original_execute(connection, statement, parameters or ())

        async def commit(connection):
            if inserted:
                raise sqlite3.OperationalError(f"injected commit failure {tmp_path}")
            return await original_commit(connection)

        with monkeypatch.context() as patch:
            patch.setattr(aiosqlite.Connection, "execute", execute)
            patch.setattr(aiosqlite.Connection, "commit", commit)
            with pytest.raises(ArtifactIntegrityError) as error:
                await store.publish(THREAD, TURN, deck[1], deck[0])
            safe_error(error.value, tmp_path)
        assert sql(tmp_path, "SELECT count(*) FROM pptx_artifacts") == [(0,)]
        paths = filesystem_matches(tmp_path, "*.pptx")
        assert len(paths) == 1
        assert paths[0].read_bytes() == deck[1]
        assert not filesystem_matches(tmp_path, ".pending-*")
        with pytest.raises(ArtifactNotFoundError):
            await store.retrieve(THREAD, TURN)
        record = await PptxArtifactStore(tmp_path, mode="offline").publish(
            THREAD, TURN, deck[1], deck[0]
        )
        assert await store.retrieve(THREAD, TURN) == (record, deck[1])

    asyncio.run(scenario())


@pytest.mark.parametrize("orphan", ["identical", "unrelated"])
def test_metadata_deletion_orphan_recovery_never_overwrites(tmp_path, deck, orphan):
    async def scenario():
        await seed(tmp_path)
        store = PptxArtifactStore(tmp_path, mode="offline")
        record = await store.publish(THREAD, TURN, deck[1], deck[0])
        path = artifact_path(tmp_path, record)
        sql(tmp_path, "DELETE FROM pptx_artifacts")
        if orphan == "unrelated":
            path.write_bytes(b"unrelated crash remnant")
            with pytest.raises(ArtifactIntegrityError):
                await store.publish(THREAD, TURN, deck[1], deck[0])
            assert path.read_bytes() == b"unrelated crash remnant"
            assert sql(tmp_path, "SELECT count(*) FROM pptx_artifacts") == [(0,)]
        else:
            assert await store.publish(THREAD, TURN, deck[1], deck[0]) == record
            sql(
                tmp_path,
                "DELETE FROM turns WHERE thread_id=? AND turn_id=?",
                (THREAD, TURN),
            )
            assert sql(tmp_path, "SELECT count(*) FROM pptx_artifacts") == [(0,)]
            with pytest.raises(ArtifactNotFoundError):
                await store.retrieve(THREAD, TURN)
            assert path.read_bytes() == deck[1]

    asyncio.run(scenario())


@pytest.mark.parametrize("value", [0, -1, True, 1.5, 10_000_001])
def test_artifact_limit_configuration_rejects_invalid_values(tmp_path, value):
    with pytest.raises(ValueError):
        PptxArtifactStore(tmp_path, mode="offline", max_bytes=value)


def test_byte_and_canonical_spec_limits_precede_publication(
    tmp_path, deck, monkeypatch
):
    async def scenario():
        await seed(tmp_path)
        spec, data = deck
        bounded = PptxArtifactStore(tmp_path, mode="offline", max_bytes=len(data) - 1)
        for invalid in (data, b"", bytearray(data), memoryview(data), "not bytes"):
            with pytest.raises(ArtifactError):
                await bounded.publish(THREAD, TURN, invalid, spec)
        assert not filesystem_matches(tmp_path, "*.pptx")
        exact = PptxArtifactStore(tmp_path, mode="offline", max_bytes=len(data))
        assert MAX_SPEC_BYTES == 1_000_000
        snapshot_size = len(canonical(spec).encode())
        # Lower the actual policy threshold to test its byte-boundary branch with
        # a valid renderable spec; the schema's total text budget is below 1 MB.
        with monkeypatch.context() as patch:
            patch.setattr(
                "mini_deerflow.pptx.artifacts.MAX_SPEC_BYTES", snapshot_size - 1
            )
            with pytest.raises(ArtifactError, match="specification limit"):
                await exact.publish(THREAD, TURN, data, spec)
        assert not filesystem_matches(tmp_path, "*.pptx")
        with monkeypatch.context() as patch:
            patch.setattr("mini_deerflow.pptx.artifacts.MAX_SPEC_BYTES", snapshot_size)
            record = await exact.publish(THREAD, TURN, data, spec)
        assert await exact.retrieve(THREAD, TURN) == (record, data)
        with monkeypatch.context() as patch:
            patch.setattr(
                "mini_deerflow.pptx.artifacts.MAX_SPEC_BYTES", snapshot_size - 1
            )
            with pytest.raises(ArtifactIntegrityError):
                await exact.retrieve(THREAD, TURN)

    asyncio.run(scenario())


@pytest.mark.parametrize(
    "path",
    [
        "../outside.pptx",
        "safe/../../outside.pptx",
        "/absolute.pptx",
        "C:/absolute.pptx",
        "C:relative.pptx",
        "\\rooted.pptx",
        "\\\\server\\share\\file.pptx",
        "\\\\?\\C:\\file.pptx",
        "safe\\file.pptx",
        "file.pptx:secret",
        "NUL",
        "con.pptx",
        "PRN",
        "AUX.txt",
        "COM1.pptx",
        "LPT9",
        "COM¹",
        "LPT².txt",
        "CONIN$",
        "CONOUT$",
        "trailing.",
        "trailing ",
        "bad?.pptx",
        "bad*.pptx",
        'bad".pptx',
        "bad|.pptx",
        "bad<.pptx",
        "bad>.pptx",
        "bad\x00.pptx",
        "",
    ],
)
def test_binary_paths_reject_traversal_windows_aliases_on_every_platform(
    tmp_path, path
):
    workspace = Workspace(tmp_path / "workspace")
    with pytest.raises(WorkspacePathError):
        workspace.write_bytes_exclusive(path, b"secret")
    with pytest.raises(WorkspacePathError):
        workspace.read_bytes(path)
    assert list(workspace.root.iterdir()) == []


def test_binary_bounds_exclusivity_and_regular_files(tmp_path):
    workspace = Workspace(tmp_path, max_read_bytes=4, max_write_bytes=4)
    data = b"\x00\xff\x80\n"
    assert workspace.write_bytes_exclusive("nested/file.pptx", data) == 4
    assert workspace.read_bytes("nested/file.pptx") == data
    with pytest.raises(FileExistsError):
        workspace.write_bytes_exclusive("nested/file.pptx", b"new")
    assert workspace.read_bytes("nested/file.pptx") == data
    with pytest.raises(WorkspaceLimitError):
        workspace.write_bytes_exclusive("too-big.pptx", data + b"x")
    (tmp_path / "large.pptx").write_bytes(data + b"x")
    with pytest.raises(WorkspaceLimitError):
        workspace.read_bytes("large.pptx")
    with pytest.raises((WorkspacePathError, OSError)):
        workspace.read_bytes("nested")
    with pytest.raises(TypeError):
        workspace.write_bytes_exclusive("not-bytes.pptx", bytearray(data))
    assert not (tmp_path / "too-big.pptx").exists()
    assert not filesystem_matches(tmp_path, ".pending-*")


def test_native_relative_path_objects_support_nested_binary_io(tmp_path):
    """Internal Path composition must work on Windows as well as POSIX."""
    workspace = Workspace(tmp_path)
    relative = Path("presentations") / "ownership-digest" / "deck.pptx"
    assert workspace.write_bytes_exclusive(relative, b"\x00\xffdeck") == 6
    assert workspace.read_bytes(relative) == b"\x00\xffdeck"
    assert not filesystem_matches(tmp_path, ".pending-*")


@pytest.mark.parametrize("operation", ["fsync", "link"])
def test_atomic_binary_failure_cleans_temporary_and_preserves_existing(
    tmp_path, monkeypatch, operation
):
    workspace = Workspace(tmp_path)
    workspace.write_bytes_exclusive("existing.pptx", b"original")

    def fail(*args, **kwargs):
        raise OSError("injected filesystem failure")

    with monkeypatch.context() as patch:
        patch.setattr(os, operation, fail)
        with pytest.raises(OSError, match="injected"):
            workspace.write_bytes_exclusive("new.pptx", b"replacement")
        with pytest.raises(OSError, match="injected"):
            workspace.write_bytes_exclusive("existing.pptx", b"replacement")
    assert not (tmp_path / "new.pptx").exists()
    assert workspace.read_bytes("existing.pptx") == b"original"
    assert not filesystem_matches(tmp_path, ".pending-*")


@pytest.mark.parametrize(
    "alias", ["file-symlink", "directory-symlink", "hardlink", "junction"]
)
def test_binary_link_aliases_are_rejected(tmp_path, alias):
    root, outside = tmp_path / "workspace", tmp_path / "outside"
    outside.mkdir()
    source = outside / "source.pptx"
    source.write_bytes(b"outside canary")
    workspace = Workspace(root)
    link = root / "alias"
    relative = "alias"
    try:
        if alias == "hardlink":
            os.link(source, link)
        elif alias == "junction":
            if os.name != "nt":
                pytest.skip("Windows junctions unavailable on this platform")
            result = subprocess.run(
                ["cmd", "/c", "mklink", "/J", str(link), str(outside)],
                capture_output=True,
                text=True,
                check=False,
            )
            if result.returncode:
                pytest.skip(f"Junction creation unavailable: {result.stderr}")
            relative = "alias/source.pptx"
        else:
            directory = alias == "directory-symlink"
            link.symlink_to(
                outside if directory else source, target_is_directory=directory
            )
            if directory:
                relative = "alias/source.pptx"
    except OSError as error:
        pytest.skip(f"{alias} unavailable: {error}")
    try:
        with pytest.raises(WorkspacePathError):
            workspace.read_bytes(relative)
        with pytest.raises((WorkspacePathError, FileExistsError)):
            workspace.write_bytes_exclusive(relative, b"overwrite")
        assert source.read_bytes() == b"outside canary"
        assert not list(outside.glob(".pending-*"))
    finally:
        if alias == "junction":
            link.rmdir()


def test_artifact_hardlink_tampering_is_not_a_verified_retry(tmp_path, deck):
    async def scenario():
        await seed(tmp_path)
        store = PptxArtifactStore(tmp_path, mode="offline")
        record = await store.publish(THREAD, TURN, deck[1], deck[0])
        path = artifact_path(tmp_path, record)
        try:
            os.link(path, _filesystem_path(tmp_path / "alias.pptx"))
        except OSError as error:
            pytest.skip(f"Hardlinks unavailable: {error}")
        with pytest.raises(ArtifactIntegrityError):
            await store.retrieve(THREAD, TURN)
        with pytest.raises(ArtifactIntegrityError):
            await store.publish(THREAD, TURN, deck[1], deck[0])
        assert path.read_bytes() == deck[1]

    asyncio.run(scenario())


@pytest.mark.parametrize("mode", ["offline", "live"])
def test_real_offline_runtime_http_download_and_fresh_app_restart(tmp_path, deck, mode):
    root = tmp_path if mode == "offline" else tmp_path / "live"

    async def prepare():
        service = DemoRuntimeService(OfflineDemoBackend(root))
        view = await service.run(RunDemoCommand(thread_id=THREAD, turn_id=TURN))
        assert view is not None
        record = await PptxArtifactStore(root, mode=mode).publish(
            THREAD, TURN, deck[1], deck[0]
        )
        return record

    record = asyncio.run(prepare())
    prefix = f"/api/{mode}/conversations/{THREAD}/turns/{TURN}"
    markdown = None
    for _ in range(2):
        # Both applications and their default services/stores are newly constructed.
        with TestClient(create_app(tmp_path)) as client:
            response = client.get(prefix + "/presentation")
            assert response.status_code == 200, response.text
            assert response.content == deck[1]
            assert digest(response.content) == record.sha256
            assert response.headers["content-type"] == PPTX_MIME
            assert (
                response.headers["content-disposition"]
                == 'attachment; filename="presentation.pptx"'
            )
            assert response.headers["x-content-type-options"] == "nosniff"
            assert response.headers["cache-control"] == "no-store"
            assert response.headers["etag"] == f'"{record.sha256}"'
            assert int(response.headers["content-length"]) == len(deck[1])
            validate_pptx(response.content, deck[0])
            assert len(Presentation(BytesIO(response.content)).slides) == 8
            report = client.get(prefix + "/artifact")
            assert report.status_code == 200
            assert report.headers["content-type"].startswith("text/markdown")
            assert (
                report.headers["content-disposition"]
                == 'attachment; filename="report.md"'
            )
            assert report.headers["x-content-type-options"] == "nosniff"
            if markdown is None:
                markdown = report.content
            assert report.content == markdown
            other_mode = "live" if mode == "offline" else "offline"
            assert (
                client.get(
                    prefix.replace(f"/{mode}/", f"/{other_mode}/") + "/presentation"
                ).status_code
                == 404
            )
            assert client.get(f"/api/{mode}/jobs").json() == []


@pytest.mark.parametrize(
    "state,expected",
    [
        ("missing", 404),
        ("failed", 404),
        ("running", 404),
        ("interrupted", 404),
        ("corrupt", 409),
        ("deleted", 409),
    ],
)
def test_http_missing_ineligible_and_corrupt_responses_are_safe(
    tmp_path, deck, state, expected
):
    async def prepare():
        repo = await seed(tmp_path, completed=False)
        store = PptxArtifactStore(tmp_path, mode="offline")
        if state != "missing":
            record = await store.publish(THREAD, TURN, deck[1], deck[0])
            if state in {"corrupt", "deleted"}:
                await complete(repo)
                path = artifact_path(tmp_path, record)
                if state == "corrupt":
                    path.write_bytes(f"private host data {tmp_path}".encode())
                else:
                    path.unlink()
            elif state == "failed":
                await repo.mark_turn_failed(
                    thread_id=THREAD, turn_id=TURN, total_tool_calls=0
                )
            elif state == "interrupted":
                await repo.mark_turn_interrupted(thread_id=THREAD, turn_id=TURN)
        else:
            await complete(repo)

    asyncio.run(prepare())
    with TestClient(create_app(tmp_path), raise_server_exceptions=False) as client:
        response = client.get(
            f"/api/offline/conversations/{THREAD}/turns/{TURN}/presentation"
        )
        assert response.status_code == expected, response.text
        assert response.json() == {
            "detail": (
                "Presentation was not found."
                if expected == 404
                else "Presentation is unavailable."
            )
        }
        safe_error(response.text, tmp_path)
        assert "private host data" not in response.text
        assert "etag" not in response.headers
        assert "content-disposition" not in response.headers


@pytest.mark.parametrize(
    "identifier",
    [
        ".hidden",
        "bad%20id",
        "C%3Asecret",
        "bad%5Cname",
        "x" * 129,
    ],
)
def test_http_path_identifiers_are_not_host_paths(tmp_path, identifier):
    with TestClient(create_app(tmp_path)) as client:
        for thread, turn in ((identifier, TURN), (THREAD, identifier)):
            response = client.get(
                f"/api/offline/conversations/{thread}/turns/{turn}/presentation"
            )
            assert response.status_code == 422
            assert response.json() == {"detail": "Invalid request."}
            safe_error(response.text, tmp_path)
    assert not filesystem_matches(tmp_path, "*.pptx")


def test_http_caller_cannot_select_a_path_or_override_metadata(tmp_path, deck):
    async def prepare():
        await seed(tmp_path)
        return await PptxArtifactStore(tmp_path, mode="offline").publish(
            THREAD, TURN, deck[1], deck[0]
        )

    record = asyncio.run(prepare())
    with TestClient(create_app(tmp_path)) as client:
        url = f"/api/offline/conversations/{THREAD}/turns/{TURN}/presentation"
        response = client.get(
            url,
            params={
                "path": str(tmp_path / "secret"),
                "sha256": "0" * 64,
                "filename": "../../secret",
                "mode": "live",
            },
        )
        assert response.status_code == 200
        assert response.content == deck[1]
        assert response.headers["etag"] == f'"{record.sha256}"'
        # Since Task 2.3 the POST presentation-generation route exists, but a
        # caller still cannot inject paths or metadata: any path-bearing body
        # fails the closed PresentationRequestDTO schema before any handler
        # logic, and generation stays unavailable without an injected planner.
        assert client.post(url, json={"path": "../../secret"}).status_code == 422
        assert client.get(url.replace(f"/{TURN}/", "/unknown/")).status_code == 404


@pytest.mark.parametrize(
    "statement",
    [
        "UPDATE pptx_artifacts SET byte_size=0",
        "UPDATE pptx_artifacts SET byte_size=10000001",
        "UPDATE pptx_artifacts SET artifact_type='html'",
        "UPDATE pptx_artifacts SET mode='unknown'",
        "UPDATE pptx_artifacts SET spec_sha256='short'",
        "UPDATE pptx_artifacts SET spec_json=printf('%1000001s','x')",
    ],
)
def test_database_rejects_invalid_metadata_independently_of_files(
    tmp_path, deck, statement
):
    async def prepare():
        await seed(tmp_path)
        store = PptxArtifactStore(tmp_path, mode="offline")
        # Initialize the real schema without needing publication to succeed.
        with pytest.raises(ArtifactNotFoundError):
            await store.retrieve(THREAD, TURN)

    asyncio.run(prepare())
    snapshot = canonical(deck[0])
    sql(
        tmp_path,
        "INSERT INTO pptx_artifacts VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        (
            "offline",
            THREAD,
            TURN,
            "pptx",
            len(deck[1]),
            digest(deck[1]),
            digest(snapshot.encode()),
            snapshot,
        ),
    )
    with pytest.raises(sqlite3.IntegrityError):
        sql(tmp_path, statement)
    assert sql(tmp_path, "SELECT byte_size, artifact_type FROM pptx_artifacts") == [
        (len(deck[1]), "pptx")
    ]


@pytest.mark.parametrize("status", ["running", "interrupted", "failed"])
def test_ineligible_download_rejects_before_any_workspace_resolution(
    tmp_path, status, monkeypatch
):
    async def scenario():
        repo = await seed(tmp_path, completed=False)
        if status == "interrupted":
            await repo.mark_turn_interrupted(thread_id=THREAD, turn_id=TURN)
        elif status == "failed":
            await repo.mark_turn_failed(
                thread_id=THREAD, turn_id=TURN, total_tool_calls=0
            )

        def forbidden(*args, **kwargs):
            pytest.fail("An ineligible download must not resolve any workspace")

        monkeypatch.setattr(SessionSandboxResolver, "resolve", forbidden)
        with pytest.raises(ArtifactNotFoundError):
            await PptxArtifactStore(tmp_path, mode="offline").retrieve(THREAD, TURN)

    asyncio.run(scenario())


@pytest.mark.parametrize("operation", ["fsync", "link"])
def test_store_atomic_filesystem_failure_is_safe_and_has_no_metadata(
    tmp_path, deck, monkeypatch, operation
):
    async def scenario():
        await seed(tmp_path)
        called = False

        def fail(*args, **kwargs):
            nonlocal called
            called = True
            raise OSError(f"injected private path {tmp_path}")

        with monkeypatch.context() as patch:
            patch.setattr(os, operation, fail)
            with pytest.raises(ArtifactIntegrityError) as error:
                await PptxArtifactStore(tmp_path, mode="offline").publish(
                    THREAD, TURN, deck[1], deck[0]
                )
            safe_error(error.value, tmp_path)
        assert called, (
            "Publication failed before reaching the intended filesystem operation"
        )
        assert sql(tmp_path, "SELECT count(*) FROM pptx_artifacts") == [(0,)]
        assert not filesystem_matches(tmp_path, "*.pptx")
        assert not filesystem_matches(tmp_path, ".pending-*")

    asyncio.run(scenario())


@pytest.mark.parametrize("ancestor", ["workspaces", "sessions", "session"])
@pytest.mark.parametrize("link_kind", ["symlink", "junction"])
def test_store_rejects_linked_sandbox_ancestors(tmp_path, deck, ancestor, link_kind):
    async def scenario():
        await seed(tmp_path)
        outside = tmp_path / "outside"
        outside.mkdir()
        canary = outside / "canary.txt"
        canary.write_bytes(b"must not change")
        session = SessionSandboxResolver(tmp_path / "workspaces").resolve(THREAD)
        link = {
            "workspaces": tmp_path / "workspaces",
            "sessions": session.parent,
            "session": session,
        }[ancestor]
        link.parent.mkdir(parents=True, exist_ok=True)
        if link_kind == "junction" and os.name != "nt":
            pytest.skip("Windows junctions unavailable on this platform")
        try:
            if link_kind == "junction":
                result = await asyncio.to_thread(
                    subprocess.run,
                    ["cmd", "/c", "mklink", "/J", str(link), str(outside)],
                    capture_output=True,
                    text=True,
                    check=False,
                )
                if result.returncode:
                    pytest.skip(f"Junction creation unavailable: {result.stderr}")
            else:
                link.symlink_to(outside, target_is_directory=True)
        except OSError as error:
            pytest.skip(f"Directory symlinks unavailable: {error}")
        try:
            with pytest.raises(ArtifactIntegrityError) as error:
                await PptxArtifactStore(tmp_path, mode="offline").publish(
                    THREAD, TURN, deck[1], deck[0]
                )
            safe_error(error.value, tmp_path)
            assert list(outside.iterdir()) == [canary]
            assert canary.read_bytes() == b"must not change"
            assert sql(tmp_path, "SELECT count(*) FROM pptx_artifacts") == [(0,)]
        finally:
            if link_kind == "junction":
                link.rmdir()

    asyncio.run(scenario())


def test_same_ids_in_both_http_modes_return_distinct_persistent_decks(
    tmp_path, deck, other_deck
):
    async def prepare():
        for mode, root, (spec, data) in (
            ("offline", tmp_path, deck),
            ("live", tmp_path / "live", other_deck),
        ):
            await DemoRuntimeService(OfflineDemoBackend(root)).run(
                RunDemoCommand(thread_id=THREAD, turn_id=TURN)
            )
            await PptxArtifactStore(root, mode=mode).publish(THREAD, TURN, data, spec)

    asyncio.run(prepare())
    for _ in range(2):
        with TestClient(create_app(tmp_path)) as client:
            responses = [
                client.get(
                    f"/api/{mode}/conversations/{THREAD}/turns/{TURN}/presentation"
                )
                for mode in ("offline", "live")
            ]
            for response, (spec, data) in zip(
                responses, (deck, other_deck), strict=True
            ):
                assert response.status_code == 200, response.text
                assert response.content == data
                assert response.headers["etag"] == f'"{digest(data)}"'
                validate_pptx(response.content, spec)
                assert len(Presentation(BytesIO(response.content)).slides) == 8
            assert responses[0].content != responses[1].content


@pytest.mark.parametrize("damage", ["non-zip", "missing-slide", "wrong-spec"])
def test_invalid_presentation_never_reaches_workspace(
    tmp_path, deck, other_deck, monkeypatch, damage
):
    async def scenario():
        await seed(tmp_path)
        spec, data = deck
        if damage == "non-zip":
            data = b"not an OOXML package"
        elif damage == "missing-slide":
            data = repack(data, remove=("ppt/slides/slide1.xml",))
        else:
            spec = other_deck[0]

        def forbidden(*args, **kwargs):
            pytest.fail("Invalid presentations must not reach workspace publication")

        monkeypatch.setattr(Workspace, "write_bytes_exclusive", forbidden)
        with pytest.raises(ArtifactIntegrityError) as error:
            await PptxArtifactStore(tmp_path, mode="offline").publish(
                THREAD, TURN, data, spec
            )
        safe_error(error.value, tmp_path)
        assert not filesystem_matches(tmp_path, "*.pptx")

    asyncio.run(scenario())
