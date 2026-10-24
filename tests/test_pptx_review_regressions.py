"""Independent regression coverage for the Issue #8 PPTX review.

Mutations preserve real packages and cached content unless that is the target
of the test. No production helpers are used to establish expected XML values.
"""

import socket
import time
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import pytest
from pptx import Presentation

from mini_deerflow.demo.phase2_pptx_offline import demo_spec
from mini_deerflow.pptx import (
    DeckSpec,
    PptxValidationError,
    render_deck,
    validate_pptx,
)

NS = {
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}
CHART = "ppt/charts/chart1.xml"
SHEET = "xl/worksheets/sheet1.xml"
WORKBOOK_RELS = "xl/_rels/workbook.xml.rels"


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("review regressions must not access the network")

    monkeypatch.setattr(socket.socket, "connect", forbidden)
    monkeypatch.setattr(socket, "create_connection", forbidden)


@pytest.fixture(scope="module")
def deck():
    spec = demo_spec()
    blob = render_deck(spec)
    assert validate_pptx(blob, spec, min_layouts=8).valid
    return spec, blob


def member_bytes(blob, member):
    with ZipFile(BytesIO(blob)) as archive:
        return archive.read(member)


def repack(blob, replacements=None, extras=None):
    output = BytesIO()
    replacements = replacements or {}
    with ZipFile(BytesIO(blob)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            target.writestr(
                entry.filename,
                replacements.get(entry.filename, source.read(entry)),
            )
        for name, data in (extras or {}).items():
            target.writestr(name, data)
    return output.getvalue()


def mutate_xml(blob, member, change):
    root = ET.fromstring(member_bytes(blob, member))
    change(root)
    return repack(blob, {member: ET.tostring(root, encoding="utf-8")})


def workbook_member(blob):
    with ZipFile(BytesIO(blob)) as archive:
        members = [
            n
            for n in archive.namelist()
            if n.startswith("ppt/embeddings/") and n.endswith(".xlsx")
        ]
    assert len(members) == 1
    return members[0]


def mutate_workbook(blob, change):
    member = workbook_member(blob)
    changed = repack(blob, {member: change(member_bytes(blob, member))})
    assert member_bytes(changed, CHART) == member_bytes(blob, CHART)
    return changed


def single_chart_payload():
    payload = demo_spec().model_dump(mode="json")
    payload["slides"] = [
        slide for slide in payload["slides"] if slide["content"]["kind"] == "chart"
    ]
    return payload


def semantic_parts(blob):
    with ZipFile(BytesIO(blob)) as archive:
        result = {}
        for name in sorted(archive.namelist()):
            data = archive.read(name)
            if name.endswith((".xml", ".rels")):
                result[name] = ET.canonicalize(data.decode("utf-8"))
            elif name.endswith(".xlsx"):
                result[name] = semantic_parts(data)
            else:
                result[name] = data
        return result


@pytest.mark.parametrize("field", ["category", "series-name"])
@pytest.mark.parametrize(
    "literal", ["=1+1", "https://example.invalid/data", "_x0041_", "_x005F_"]
)
def test_chart_labels_remain_literal_in_cache_and_editable_workbook(field, literal):
    payload = single_chart_payload()
    chart = payload["slides"][0]["content"]
    if field == "category":
        chart["categories"][0] = literal
        address = "A2"
        cache_path = "c:cat/c:strRef/c:strCache/c:pt/c:v"
    else:
        chart["series"][0]["name"] = literal
        address = "B1"
        cache_path = "c:tx/c:strRef/c:strCache/c:pt/c:v"
    spec = DeckSpec.model_validate(payload)
    blob = render_deck(spec)
    assert validate_pptx(blob, spec).valid
    root = ET.fromstring(member_bytes(blob, CHART))
    assert root.find(".//c:ser", NS).find(cache_path, NS).text == literal
    workbook = member_bytes(blob, workbook_member(blob))
    sheet = ET.fromstring(member_bytes(workbook, SHEET))
    cell = sheet.find(f".//s:c[@r='{address}']", NS)
    assert cell is not None and cell.get("t") == "s"
    index = int(cell.find("s:v", NS).text)
    strings = ET.fromstring(member_bytes(workbook, "xl/sharedStrings.xml"))
    encoded = "".join(strings[index].itertext())
    assert encoded == literal.replace("_x", "_x005F_x")
    with ZipFile(BytesIO(workbook)) as archive:
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                part = ET.fromstring(archive.read(name))
                assert not part.findall(".//s:f", NS)
                if name.endswith(".rels"):
                    assert all(r.get("TargetMode") != "External" for r in part)


def test_semantic_determinism_includes_workbook_across_wallclock_gap():
    spec = demo_spec()
    first = render_deck(spec)
    time.sleep(1.1)
    second = render_deck(spec)
    assert validate_pptx(first, spec, min_layouts=8).valid
    assert validate_pptx(second, spec, min_layouts=8).valid
    assert semantic_parts(first) == semantic_parts(second)


@pytest.mark.parametrize(
    "raw,canonical",
    [
        (1.2345678901234567, 1.23456789012346),
        (0.12345678901234567, 0.123456789012346),
        (-1.2345678901234567, -1.23456789012346),
        (999999999.9999999, 1_000_000_000.0),
    ],
)
def test_precision_policy_is_canonical_in_spec_cache_workbook_and_table(raw, canonical):
    payload = demo_spec().model_dump(mode="json")
    payload["slides"][5]["content"]["series"][0]["values"][0] = raw
    payload["slides"][6]["content"]["rows"][0][1] = raw
    spec = DeckSpec.model_validate(payload)
    assert spec.slides[5].content.series[0].values[0] == canonical
    assert spec.slides[6].content.rows[0][1] == canonical
    assert DeckSpec.model_validate_json(spec.model_dump_json()) == spec
    blob = render_deck(spec)
    assert validate_pptx(blob, spec, min_layouts=8).valid
    reopened = Presentation(BytesIO(blob))
    chart = next(sh.chart for sh in reopened.slides[5].shapes if sh.has_chart)
    assert chart.series[0].values[0] == canonical
    table = next(sh.table for sh in reopened.slides[6].shapes if sh.has_table)
    assert table.cell(1, 1).text == str(canonical)
    workbook = member_bytes(blob, workbook_member(blob))
    sheet = ET.fromstring(member_bytes(workbook, SHEET))
    assert float(sheet.find(".//s:c[@r='B2']/s:v", NS).text) == canonical


def test_workbook_formula_rejected_even_when_cached_value_matches(deck):
    spec, blob = deck

    def change(root):
        cell = root.find(".//s:c[@r='B2']", NS)
        assert cell is not None and cell.find("s:v", NS).text == "12"
        formula = ET.Element(f"{{{NS['s']}}}f")
        formula.text = "99+99"
        cell.insert(0, formula)
        assert cell.find("s:v", NS).text == "12"

    bad = mutate_workbook(blob, lambda wb: mutate_xml(wb, SHEET, change))
    with pytest.raises(PptxValidationError, match="formulas"):
        validate_pptx(bad, spec)


@pytest.mark.parametrize("series_index", [0, 1])
@pytest.mark.parametrize(
    "reference", ["c:tx/c:strRef/c:f", "c:cat/c:strRef/c:f", "c:val/c:numRef/c:f"]
)
def test_chart_references_must_match_each_series_workbook_ranges(
    deck, series_index, reference
):
    spec, blob = deck

    def change(root):
        node = root.findall(".//c:ser", NS)[series_index].find(reference, NS)
        assert node is not None
        node.text = "Sheet1!$Z$100:$Z$102"

    bad = mutate_xml(blob, CHART, change)
    assert member_bytes(bad, workbook_member(bad)) == member_bytes(
        blob, workbook_member(blob)
    )
    with pytest.raises(PptxValidationError, match="references"):
        validate_pptx(bad, spec)


def test_nonfirst_series_category_cache_is_checked(deck):
    spec, blob = deck

    def change(root):
        category = root.findall(".//c:ser", NS)[1].find(
            "c:cat/c:strRef/c:strCache/c:pt/c:v", NS
        )
        assert category is not None
        category.text = "changed category"

    with pytest.raises(PptxValidationError, match="category cache"):
        validate_pptx(mutate_xml(blob, CHART, change), spec)


@pytest.mark.parametrize(
    "slide_index,name", [(0, "subtitle"), (5, "chart"), (6, "table")]
)
@pytest.mark.parametrize("axis", ["width", "height"])
def test_tiny_positive_native_object_geometry_is_rejected(
    deck, slide_index, name, axis
):
    spec, blob = deck
    prs = Presentation(BytesIO(blob))
    shape = next(sh for sh in prs.slides[slide_index].shapes if sh.name == name)
    setattr(shape, axis, 1)
    output = BytesIO()
    prs.save(output)
    with pytest.raises(PptxValidationError, match="geometry"):
        validate_pptx(output.getvalue(), spec)


@pytest.mark.parametrize("index", [-1, 999999])
def test_invalid_workbook_shared_string_indices_raise_public_error(deck, index):
    spec, blob = deck

    def change(root):
        value = root.find(".//s:c[@r='A2']/s:v", NS)
        assert value is not None
        value.text = str(index)

    bad = mutate_workbook(blob, lambda wb: mutate_xml(wb, SHEET, change))
    with pytest.raises(PptxValidationError, match="shared-string index"):
        validate_pptx(bad, spec)


def test_external_relationship_inside_embedded_workbook_is_rejected(deck):
    spec, blob = deck

    def change(root):
        ET.SubElement(
            root,
            f"{{{NS['r']}}}Relationship",
            {
                "Id": "reviewExternal",
                "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                "Target": "https://example.invalid/data",
                "TargetMode": "External",
            },
        )

    bad = mutate_workbook(blob, lambda wb: mutate_xml(wb, WORKBOOK_RELS, change))
    with pytest.raises(PptxValidationError, match="external relationships"):
        validate_pptx(bad, spec)


def test_original_zip_filename_is_checked_before_platform_normalization(deck):
    spec, blob = deck
    name = "bad\\path.xml"
    entry = ZipInfo(name)
    entry.filename = name  # Preserve the literal backslash on Windows, too.
    bad = repack(blob, extras={entry: b"<root/>"})
    with ZipFile(BytesIO(bad)) as archive:
        assert name in [info.orig_filename for info in archive.infolist()]
    with pytest.raises(PptxValidationError, match="unsafe ZIP member path"):
        validate_pptx(bad, spec)


def test_declared_binary_macro_project_in_embedded_workbook_is_rejected(deck):
    spec, blob = deck

    def change(workbook):
        def relationship(root):
            ET.SubElement(
                root,
                f"{{{NS['r']}}}Relationship",
                {
                    "Id": "reviewMacro",
                    "Type": "http://schemas.microsoft.com/office/2006/relationships/vbaProject",
                    "Target": "vbaProject.bin",
                },
            )

        def content_type(root):
            ET.SubElement(
                root,
                f"{{{NS['ct']}}}Override",
                {
                    "PartName": "/xl/vbaProject.bin",
                    "ContentType": "application/vnd.ms-office.vbaProject",
                },
            )

        workbook = mutate_xml(workbook, WORKBOOK_RELS, relationship)
        workbook = mutate_xml(workbook, "[Content_Types].xml", content_type)
        # VBA is opaque binary here; rejection must not depend on executing it.
        return repack(workbook, extras={"xl/vbaProject.bin": b"opaque VBA project"})

    with pytest.raises(PptxValidationError):
        validate_pptx(mutate_workbook(blob, change), spec)


def test_nonfirst_series_category_point_indices_must_be_unique(deck):
    spec, blob = deck

    def change(root):
        points = root.findall(".//c:ser", NS)[1].findall(
            "c:cat/c:strRef/c:strCache/c:pt", NS
        )
        assert len(points) == 3
        for point in points:
            point.set("idx", "0")

    with pytest.raises(PptxValidationError):
        validate_pptx(mutate_xml(blob, CHART, change), spec)
